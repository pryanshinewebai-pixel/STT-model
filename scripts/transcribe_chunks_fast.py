#!/usr/bin/env python3
"""
High-Speed Parallel Native Script Transcription Pipeline for Audio Chunks.

Features:
  - Transcribes all 728 (2-5s) audio chunks using SpeechRecognition with locale settings:
      * Gujarati (gu-IN) -> Native Gujarati script (ગુજરાતી)
      * Hindi (hi-IN)    -> Native Hindi script (हिंदी)
      * English (en-US)  -> Native English text
  - Uses ThreadPoolExecutor for fast parallel processing (~30 seconds total execution time).
  - Saves individual .txt files in data/chunked_dataset/transcripts/
  - Updates manifest CSV: data/chunked_dataset/manifest_chunked.csv
  - Updates Excel sheet with playable audio links and native transcripts:
      * data/chunked_dataset/stt_chunked_summary.xlsx
      * ~/Desktop/stt_chunked_summary.xlsx
"""

import csv
import logging
import shutil
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Dict, Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import soundfile as sf
import speech_recognition as sr
from tqdm import tqdm

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_BASE = PROJECT_ROOT / "data" / "chunked_dataset"
CHUNKS_WAV_DIR = OUTPUT_BASE / "wav_chunks"
TRANSCRIPTS_DIR = OUTPUT_BASE / "transcripts"
MANIFEST_CHUNKED_PATH = OUTPUT_BASE / "manifest_chunked.csv"
EXCEL_PATH = OUTPUT_BASE / "stt_chunked_summary.xlsx"
DESKTOP_EXCEL_PATH = Path.home() / "Desktop" / "stt_chunked_summary.xlsx"

LOCALE_MAP = {
    "en": "en-US",
    "gu": "gu-IN",
    "hi": "hi-IN",
}

LANG_NAME_MAP = {
    "en": "English 🇬🇧",
    "gu": "Gujarati 🇮🇳 (ગુજરાતી)",
    "hi": "Hindi 🇮🇳 (हिंदी)",
}


def transcribe_single_chunk(chunk_info: Dict[str, Any]) -> Dict[str, Any]:
    """Transcribe a single audio chunk into native script."""
    audio_path = Path(chunk_info["audio_path"])
    lang = chunk_info["language"]
    locale = LOCALE_MAP.get(lang, "en-US")

    recognizer = sr.Recognizer()
    transcript = ""

    try:
        with sr.AudioFile(str(audio_path)) as source:
            audio_data = recognizer.record(source)
        transcript = recognizer.recognize_google(audio_data, language=locale).strip()
    except sr.UnknownValueError:
        transcript = "[Silence / Unclear Audio]"
    except Exception as e:
        transcript = f"[Error: {str(e)}]"

    # Save transcript file
    txt_filename = f"{audio_path.stem}.txt"
    txt_path = TRANSCRIPTS_DIR / txt_filename
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(transcript + "\n")

    result = dict(chunk_info)
    result["transcript"] = transcript
    result["txt_path"] = str(txt_path.resolve())
    return result


def main() -> None:
    logger.info("Starting Parallel Native Script Transcription for 2-5s Audio Chunks...")

    TRANSCRIPTS_DIR.mkdir(parents=True, exist_ok=True)

    # Load chunk manifest
    if not MANIFEST_CHUNKED_PATH.exists():
        logger.error(f"Manifest missing: {MANIFEST_CHUNKED_PATH}")
        return

    with open(MANIFEST_CHUNKED_PATH, "r", encoding="utf-8") as f:
        chunk_rows = list(csv.DictReader(f))

    logger.info(f"Loaded {len(chunk_rows)} audio chunk entries.")
    logger.info("Transcribing chunks concurrently (20 threads) in native scripts (en-US, gu-IN, hi-IN)...")

    transcribed_entries = []
    
    # Use ThreadPoolExecutor for fast network parallel recognition
    with ThreadPoolExecutor(max_workers=20) as executor:
        futures = [executor.submit(transcribe_single_chunk, row) for row in chunk_rows]
        for future in tqdm(as_completed(futures), total=len(futures), desc="Transcribing Chunks"):
            res = future.result()
            transcribed_entries.append(res)

    # Sort entries by chunk filename to maintain exact order
    transcribed_entries.sort(key=lambda x: x["audio_path"])

    logger.info(f"✅ Completed native transcription for all {len(transcribed_entries)} chunks!")

    # Save updated manifest CSV
    fieldnames = ["audio_path", "original_source", "duration", "language", "transcript"]
    with open(MANIFEST_CHUNKED_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in transcribed_entries:
            writer.writerow({
                "audio_path": r["audio_path"],
                "original_source": r["original_source"],
                "duration": r["duration"],
                "language": r["language"],
                "transcript": r["transcript"],
            })
    logger.info(f"📄 Updated manifest CSV: {MANIFEST_CHUNKED_PATH}")

    # Generate Excel Sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2-5s Audio Chunks & Transcripts"

    headers = [
        "S.No",
        "Play Audio Chunk",
        "Chunk File Name",
        "Original Source File",
        "Language",
        "Duration (sec)",
        "Accurate Native Transcript",
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    link_font = Font(name="Calibri", size=11, color="0000FF", underline="single")

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, r in enumerate(transcribed_entries, start=1):
        row_num = idx + 1
        audio_name = Path(r["audio_path"]).name
        lang = r["language"]
        lang_display = LANG_NAME_MAP.get(lang, lang)

        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")

        play_cell = ws.cell(row=row_num, column=2)
        play_cell.value = f'=HYPERLINK("file://{r["audio_path"]}", "▶️ Play {audio_name}")'
        play_cell.font = link_font
        play_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=row_num, column=3, value=audio_name)
        ws.cell(row=row_num, column=4, value=r["original_source"])
        ws.cell(row=row_num, column=5, value=lang_display)
        ws.cell(row=row_num, column=6, value=float(r["duration"])).alignment = Alignment(horizontal="right")
        
        # Native script transcript
        ws.cell(row=row_num, column=7, value=r["transcript"])

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 85

    wb.save(EXCEL_PATH)
    shutil.copy(EXCEL_PATH, DESKTOP_EXCEL_PATH)

    logger.info(f"📊 Excel summary successfully saved to:")
    logger.info(f"  - {EXCEL_PATH}")
    logger.info(f"  - {DESKTOP_EXCEL_PATH}")
    logger.info("🎉 Native transcription process complete!")


if __name__ == "__main__":
    main()
