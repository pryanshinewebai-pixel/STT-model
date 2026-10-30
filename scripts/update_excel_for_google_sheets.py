#!/usr/bin/env python3
"""
Generate Google Sheets Compatible Excel File & Web Audio Dashboard with Live Editable Transcripts.

Features:
  - Generates Excel sheet formatted specifically for Google Sheets compatibility (http://localhost:8000/wav_chunks/...).
  - Creates interactive HTML web player with Auto-Play & Live Manual Text Editing at data/chunked_dataset/index.html
"""

import csv
import json
import logging
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_BASE = PROJECT_ROOT / "data" / "chunked_dataset"
MANIFEST_CHUNKED = OUTPUT_BASE / "manifest_chunked.csv"
MANIFEST_TOKENIZED = OUTPUT_BASE / "manifest_tokenized.csv"
EXCEL_PATH = OUTPUT_BASE / "stt_chunked_summary.xlsx"
DESKTOP_EXCEL_PATH = Path.home() / "Desktop" / "stt_chunked_summary.xlsx"
HTML_PLAYER_PATH = OUTPUT_BASE / "index.html"

LANG_NAME_MAP = {
    "en": "English 🇬🇧",
    "gu": "Gujarati 🇮🇳 (ગુજરાતી)",
    "hi": "Hindi 🇮🇳 (हिंदी)",
}

PORT = 8000


def main() -> None:
    logger.info("Updating Excel File & Generating Web Dashboard with Live Editable Transcripts...")

    if not MANIFEST_TOKENIZED.exists():
        manifest_file = MANIFEST_CHUNKED
    else:
        manifest_file = MANIFEST_TOKENIZED

    if not manifest_file.exists():
        logger.error(f"Manifest missing: {manifest_file}")
        return

    with open(manifest_file, "r", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    logger.info(f"Loaded {len(rows)} audio chunk entries.")
    rows.sort(key=lambda x: x["audio_path"])

    # Update Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2-5s Chunks & Transcripts"

    headers = [
        "S.No",
        "Play in Google Sheets",
        "Play Local File",
        "Chunk File Name",
        "Language",
        "Duration (sec)",
        "Spoken Audio Transcript (Editable)",
        "Token Count",
        "Character CTC Token IDs",
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    link_font_gs = Font(name="Calibri", size=11, color="008000", underline="single", bold=True)
    link_font_local = Font(name="Calibri", size=11, color="0000FF", underline="single")

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, r in enumerate(rows, start=1):
        row_num = idx + 1
        audio_path = r["audio_path"]
        audio_name = Path(audio_path).name
        lang = r["language"]
        lang_display = LANG_NAME_MAP.get(lang, lang)
        duration = float(r["duration"])
        transcript = r["transcript"]
        num_tokens = r.get("num_tokens", "")
        token_ids_str = r.get("token_ids_str", "")

        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")

        gs_play_cell = ws.cell(row=row_num, column=2)
        gs_url = f"http://localhost:{PORT}/wav_chunks/{audio_name}"
        gs_play_cell.value = f'=HYPERLINK("{gs_url}", "▶️ Play in Browser / Sheets")'
        gs_play_cell.font = link_font_gs
        gs_play_cell.alignment = Alignment(horizontal="center")

        local_play_cell = ws.cell(row=row_num, column=3)
        local_url = f"file://{audio_path}"
        local_play_cell.value = f'=HYPERLINK("{local_url}", "📁 Play Local File")'
        local_play_cell.font = link_font_local
        local_play_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=row_num, column=4, value=audio_name)
        ws.cell(row=row_num, column=5, value=lang_display)
        ws.cell(row=row_num, column=6, value=duration).alignment = Alignment(horizontal="right")
        ws.cell(row=row_num, column=7, value=transcript)
        ws.cell(row=row_num, column=8, value=num_tokens).alignment = Alignment(horizontal="right")
        ws.cell(row=row_num, column=9, value=token_ids_str)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 75
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 65

    wb.save(EXCEL_PATH)
    shutil.copy(EXCEL_PATH, DESKTOP_EXCEL_PATH)

    # Generate HTML Player with Auto-Play & Live Editing
    html_content = generate_html_player(rows)
    with open(HTML_PLAYER_PATH, "w", encoding="utf-8") as f:
        f.write(html_content)

    logger.info(f"🌐 Saved Editable Web Dashboard to: {HTML_PLAYER_PATH}")


def generate_html_player(rows) -> str:
    """Generate responsive HTML audio dashboard with live editable transcripts."""
    data_json = json.dumps(rows, ensure_ascii=False)

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>STT Audio Chunks Dashboard - Live Editable Transcripts</title>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;600;700&display=swap" rel="stylesheet">
    <style>
        * {{ box-sizing: border-box; margin: 0; padding: 0; font-family: 'Inter', sans-serif; }}
        body {{ background: #0f172a; color: #f8fafc; padding: 20px; padding-bottom: 140px; }}
        .container {{ max-width: 1400px; margin: 0 auto; }}
        header {{ text-align: center; margin-bottom: 25px; padding: 22px; background: #1e293b; border-radius: 14px; box-shadow: 0 10px 25px rgba(0,0,0,0.3); border: 1px solid #334155; }}
        h1 {{ font-size: 24px; color: #38bdf8; margin-bottom: 8px; font-weight: 700; }}
        p {{ color: #94a3b8; font-size: 14px; }}

        .controls {{ display: flex; gap: 12px; margin-bottom: 20px; flex-wrap: wrap; align-items: center; }}
        input.search-input, select {{ padding: 12px 16px; border-radius: 10px; border: 1px solid #334155; background: #1e293b; color: #fff; font-size: 14px; outline: none; transition: 0.2s; }}
        input.search-input {{ flex: 2; min-width: 250px; }}
        select {{ flex: 1; min-width: 180px; }}
        input:focus, select:focus {{ border-color: #38bdf8; box-shadow: 0 0 0 3px rgba(56, 189, 248, 0.2); }}

        .btn {{ padding: 12px 20px; border-radius: 10px; border: none; font-weight: 600; font-size: 14px; cursor: pointer; display: inline-flex; align-items: center; gap: 8px; transition: all 0.2s ease; }}
        .btn-autoplay {{ background: #10b981; color: white; box-shadow: 0 4px 14px rgba(16, 185, 129, 0.4); }}
        .btn-autoplay:hover {{ background: #059669; transform: translateY(-1px); }}
        .btn-autoplay.active {{ background: #ef4444; box-shadow: 0 4px 14px rgba(239, 68, 68, 0.4); }}

        .stats {{ background: #1e293b; padding: 12px 18px; border-radius: 10px; font-size: 14px; font-weight: 600; color: #38bdf8; border: 1px solid #334155; }}

        table {{ width: 100%; border-collapse: collapse; background: #1e293b; border-radius: 14px; overflow: hidden; box-shadow: 0 10px 25px rgba(0,0,0,0.3); border: 1px solid #334155; }}
        th, td {{ padding: 12px 14px; text-align: left; border-bottom: 1px solid #334155; font-size: 14px; vertical-align: middle; }}
        th {{ background: #0f172a; color: #94a3b8; text-transform: uppercase; font-size: 12px; letter-spacing: 0.5px; font-weight: 700; }}
        tr {{ transition: background 0.15s; }}
        tr:hover {{ background: #26334d; }}
        
        tr.playing-row {{ background: #1e3a8a !important; border-left: 4px solid #38bdf8; }}
        tr.playing-row td {{ color: #ffffff; font-weight: 500; }}

        .play-row-btn {{ background: #0284c7; color: white; border: none; padding: 8px 14px; border-radius: 8px; cursor: pointer; font-weight: 600; font-size: 13px; transition: 0.2s; white-space: nowrap; }}
        .play-row-btn:hover {{ background: #0369a1; transform: scale(1.03); }}

        /* Editable Transcript Box */
        .edit-container {{ display: flex; gap: 8px; align-items: center; width: 100%; }}
        .edit-input {{
            flex: 1;
            padding: 8px 12px;
            border-radius: 8px;
            border: 1px solid #475569;
            background: #0f172a;
            color: #f8fafc;
            font-size: 15px;
            font-weight: 500;
            outline: none;
            transition: all 0.2s;
        }}
        .edit-input:focus {{
            border-color: #38bdf8;
            background: #1e293b;
            box-shadow: 0 0 0 3px rgba(56, 189, 248, 0.25);
        }}
        .edit-input.modified {{
            border-color: #f59e0b;
            background: #281d0c;
        }}

        .btn-save {{
            background: #334155;
            color: #94a3b8;
            border: 1px solid #475569;
            padding: 8px 12px;
            border-radius: 8px;
            font-size: 12px;
            font-weight: 600;
            cursor: pointer;
            white-space: nowrap;
            transition: 0.2s;
        }}
        .btn-save.active {{
            background: #f59e0b;
            color: #000;
            font-weight: 700;
            border-color: #fbbf24;
            box-shadow: 0 0 10px rgba(245, 158, 11, 0.4);
        }}
        .btn-save.saved {{
            background: #10b981 !important;
            color: white !important;
            border-color: #34d399 !important;
        }}

        .badge {{ padding: 4px 10px; border-radius: 20px; font-size: 12px; font-weight: 600; display: inline-block; white-space: nowrap; }}
        .badge-en {{ background: #1e3a8a; color: #93c5fd; }}
        .badge-gu {{ background: #831843; color: #fbcfe8; }}
        .badge-hi {{ background: #14532d; color: #bbf7d0; }}

        audio.inline-player {{ height: 32px; width: 170px; outline: none; border-radius: 6px; }}

        /* Sticky Bottom Player */
        .now-playing-bar {{
            position: fixed;
            bottom: 0;
            left: 0;
            right: 0;
            background: #0f172a;
            border-top: 2px solid #38bdf8;
            padding: 16px 24px;
            box-shadow: 0 -10px 30px rgba(0,0,0,0.5);
            display: flex;
            align-items: center;
            justify-content: space-between;
            gap: 20px;
            z-index: 999;
            backdrop-filter: blur(10px);
        }}
        .np-info {{ flex: 1; overflow: hidden; }}
        .np-title {{ font-size: 14px; font-weight: 700; color: #38bdf8; margin-bottom: 4px; display: flex; align-items: center; gap: 10px; }}
        .np-transcript {{ font-size: 15px; color: #f8fafc; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .np-controls {{ display: flex; align-items: center; gap: 12px; }}
        .np-controls button {{ background: #1e293b; color: white; border: 1px solid #334155; padding: 10px 14px; border-radius: 8px; cursor: pointer; font-weight: 600; font-size: 13px; }}
        .np-controls button:hover {{ background: #334155; border-color: #38bdf8; }}
        audio.global-player {{ outline: none; border-radius: 8px; max-width: 340px; width: 100%; height: 36px; }}

        /* Toast notification */
        .toast {{
            position: fixed;
            top: 20px;
            right: 20px;
            background: #10b981;
            color: white;
            padding: 12px 20px;
            border-radius: 10px;
            font-weight: 600;
            box-shadow: 0 10px 25px rgba(0,0,0,0.3);
            display: none;
            z-index: 10000;
            animation: fadeIn 0.3s ease;
        }}
        @keyframes fadeIn {{ from {{ opacity: 0; transform: translateY(-10px); }} to {{ opacity: 1; transform: translateY(0); }} }}
    </style>
</head>
<body>
    <div class="toast" id="toastMsg">✅ Transcript saved successfully!</div>

    <div class="container">
        <header>
            <h1>🎙️ STT Audio Chunks Dashboard - Live Manual Editing</h1>
            <p>Edit transcripts directly in the table below and press <strong>💾 Save</strong> to update local files & Excel instantly!</p>
        </header>

        <div class="controls">
            <input type="text" class="search-input" id="searchInput" placeholder="🔍 Search transcript or filename..." oninput="filterData()">
            <select id="langFilter" onchange="filterData()">
                <option value="all">All Languages</option>
                <option value="gu">Gujarati 🇮🇳 (ગુજરાતી)</option>
                <option value="hi">Hindi 🇮🇳 (हिंदी)</option>
                <option value="en">English 🇬🇧</option>
            </select>
            <button class="btn btn-autoplay" id="autoPlayBtn" onclick="toggleAutoPlay()">
                <span id="autoPlayIcon">▶️</span> <span id="autoPlayText">Auto-Play All Chunks</span>
            </button>
            <div class="stats" id="statsCount">Total Chunks: 0</div>
        </div>

        <table>
            <thead>
                <tr>
                    <th style="width: 45px;">#</th>
                    <th style="width: 85px;">Action</th>
                    <th style="width: 180px;">Audio Player</th>
                    <th style="width: 220px;">Chunk File Name</th>
                    <th style="width: 130px;">Language</th>
                    <th style="width: 75px;">Dur</th>
                    <th>Manually Editable Spoken Transcript</th>
                </tr>
            </thead>
            <tbody id="tableBody"></tbody>
        </table>
    </div>

    <!-- Sticky Now Playing Bottom Bar -->
    <div class="now-playing-bar" id="nowPlayingBar">
        <div class="np-info">
            <div class="np-title">
                <span id="npFileName">🔊 Select Chunk to Play</span>
                <span id="npBadge" class="badge badge-en">Ready</span>
            </div>
            <div class="np-transcript" id="npTranscript">Click any "▶️ Play" button or press "Auto-Play All" to listen continuously.</div>
        </div>
        <div class="np-controls">
            <button onclick="playPrevTrack()">⏮️ Prev</button>
            <audio id="globalAudioPlayer" class="global-player" controls preload="auto"></audio>
            <button onclick="playNextTrack()">⏭️ Next</button>
        </div>
    </div>

    <script>
        const dataset = {data_json};
        const langMap = {{
            'en': 'English 🇬🇧',
            'gu': 'Gujarati 🇮🇳 (ગુજરાતી)',
            'hi': 'Hindi 🇮🇳 (हिंदी)'
        }};

        let filteredData = [...dataset];
        let currentIndex = -1;
        let isAutoPlay = false;

        const globalAudio = document.getElementById('globalAudioPlayer');
        const autoPlayBtn = document.getElementById('autoPlayBtn');
        const autoPlayText = document.getElementById('autoPlayText');
        const autoPlayIcon = document.getElementById('autoPlayIcon');

        function renderTable() {{
            const tbody = document.getElementById('tableBody');
            tbody.innerHTML = '';
            document.getElementById('statsCount').innerText = `Total Chunks: ${{filteredData.length}}`;

            filteredData.forEach((item, index) => {{
                const audioName = item.audio_path.split('/').pop();
                const audioUrl = `/wav_chunks/${{audioName}}`;
                const langClass = `badge-${{item.language}}`;
                const langLabel = langMap[item.language] || item.language;

                const tr = document.createElement('tr');
                tr.id = `row-${{index}}`;
                if (index === currentIndex) tr.classList.add('playing-row');

                tr.innerHTML = `
                    <td>${{index + 1}}</td>
                    <td>
                        <button class="play-row-btn" onclick="playTrackByIndex(${{index}})">▶️ Play</button>
                    </td>
                    <td>
                        <audio class="inline-player" controls src="${{audioUrl}}" preload="none"></audio>
                    </td>
                    <td style="font-weight: 600; color: #38bdf8;">${{audioName}}</td>
                    <td><span class="badge ${{langClass}}">${{langLabel}}</span></td>
                    <td style="text-align: right;">${{item.duration}}s</td>
                    <td>
                        <div class="edit-container">
                            <input type="text" class="edit-input" id="input-${{index}}" 
                                   value="${{escapeHtml(item.transcript)}}" 
                                   oninput="onInputChange(${{index}})" 
                                   onkeydown="onInputKeyDown(event, ${{index}})">
                            <button class="btn-save" id="saveBtn-${{index}}" onclick="saveTranscript(${{index}})">💾 Save</button>
                        </div>
                    </td>
                `;
                tbody.appendChild(tr);
            }});
        }}

        function escapeHtml(str) {{
            return (str || '').replace(/&/g, "&amp;").replace(/</g, "&lt;").replace(/>/g, "&gt;").replace(/"/g, "&quot;");
        }}

        function onInputChange(index) {{
            const input = document.getElementById(`input-${{index}}`);
            const saveBtn = document.getElementById(`saveBtn-${{index}}`);
            if (!input || !saveBtn) return;

            const orig = filteredData[index].transcript;
            if (input.value !== orig) {{
                input.classList.add('modified');
                saveBtn.classList.add('active');
                saveBtn.innerText = '💾 Save';
            }} else {{
                input.classList.remove('modified');
                saveBtn.classList.remove('active');
            }}
        }}

        function onInputKeyDown(event, index) {{
            if (event.key === 'Enter') {{
                event.preventDefault();
                saveTranscript(index);
            }}
        }}

        async function saveTranscript(index) {{
            const input = document.getElementById(`input-${{index}}`);
            const saveBtn = document.getElementById(`saveBtn-${{index}}`);
            if (!input || !saveBtn) return;

            const newText = input.value.trim();
            const item = filteredData[index];
            const audioName = item.audio_path.split('/').pop();

            saveBtn.innerText = '⌛ Saving...';

            try {{
                const res = await fetch('/api/save_transcript', {{
                    method: 'POST',
                    headers: {{ 'Content-Type': 'application/json' }},
                    body: JSON.stringify({{ audio_name: audioName, transcript: newText }})
                }});

                const data = await res.json();
                if (data.status === 'success') {{
                    // Update dataset state
                    item.transcript = newText;

                    input.classList.remove('modified');
                    saveBtn.classList.remove('active');
                    saveBtn.classList.add('saved');
                    saveBtn.innerText = '✅ Saved!';

                    showToast(`✅ Saved '${{audioName}}' successfully!`);

                    setTimeout(() => {{
                        saveBtn.classList.remove('saved');
                        saveBtn.innerText = '💾 Save';
                    }}, 2000);
                }} else {{
                    alert('Error saving: ' + data.message);
                    saveBtn.innerText = '💾 Save';
                }}
            }} catch (err) {{
                alert('Connection error saving transcript: ' + err);
                saveBtn.innerText = '💾 Save';
            }}
        }}

        function showToast(msg) {{
            const toast = document.getElementById('toastMsg');
            if (!toast) return;
            toast.innerText = msg;
            toast.style.display = 'block';
            setTimeout(() => {{ toast.style.display = 'none'; }}, 3000);
        }}

        function filterData() {{
            const query = document.getElementById('searchInput').value.toLowerCase();
            const lang = document.getElementById('langFilter').value;

            filteredData = dataset.filter(item => {{
                const matchesLang = (lang === 'all' || item.language === lang);
                const matchesSearch = item.transcript.toLowerCase().includes(query) || item.audio_path.toLowerCase().includes(query);
                return matchesLang && matchesSearch;
            }});

            currentIndex = -1;
            renderTable();
        }}

        function playTrackByIndex(index) {{
            if (index < 0 || index >= filteredData.length) return;

            if (currentIndex >= 0) {{
                const oldRow = document.getElementById(`row-${{currentIndex}}`);
                if (oldRow) oldRow.classList.remove('playing-row');
            }}

            currentIndex = index;
            const item = filteredData[currentIndex];
            const audioName = item.audio_path.split('/').pop();
            const audioUrl = `/wav_chunks/${{audioName}}`;

            const newRow = document.getElementById(`row-${{currentIndex}}`);
            if (newRow) {{
                newRow.classList.add('playing-row');
                newRow.scrollIntoView({{ behavior: 'smooth', block: 'center' }});
            }}

            const npFileName = document.getElementById('npFileName');
            const npBadge = document.getElementById('npBadge');
            const npTranscript = document.getElementById('npTranscript');

            if (npFileName) npFileName.innerText = `🔊 Playing [${{currentIndex + 1}} / ${{filteredData.length}}] - ${{audioName}}`;
            if (npBadge) {{
                npBadge.className = `badge badge-${{item.language}}`;
                npBadge.innerText = langMap[item.language] || item.language;
            }}
            if (npTranscript) npTranscript.innerText = `"${{item.transcript}}"`;

            globalAudio.src = audioUrl;
            globalAudio.load();
            const playPromise = globalAudio.play();
            if (playPromise !== undefined) {{
                playPromise.catch(e => console.log('Audio playback info:', e));
            }}
        }}

        function playNextTrack() {{
            if (currentIndex + 1 < filteredData.length) {{
                playTrackByIndex(currentIndex + 1);
            }} else if (isAutoPlay) {{
                playTrackByIndex(0);
            }}
        }}

        function playPrevTrack() {{
            if (currentIndex > 0) {{
                playTrackByIndex(currentIndex - 1);
            }}
        }}

        function toggleAutoPlay() {{
            isAutoPlay = !isAutoPlay;
            if (isAutoPlay) {{
                autoPlayBtn.classList.add('active');
                autoPlayIcon.innerText = '⏸️';
                autoPlayText.innerText = 'Pause Auto-Play';
                if (currentIndex === -1) {{
                    playTrackByIndex(0);
                }} else {{
                    globalAudio.play().catch(e => console.log('Play defer:', e));
                }}
            }} else {{
                autoPlayBtn.classList.remove('active');
                autoPlayIcon.innerText = '▶️';
                autoPlayText.innerText = 'Auto-Play All Chunks';
                globalAudio.pause();
            }}
        }}

        globalAudio.onended = () => {{
            if (isAutoPlay) {{
                playNextTrack();
            }}
        }};

        // Initialize table
        filterData();
    </script>
</body>
</html>
"""


if __name__ == "__main__":
    main()
