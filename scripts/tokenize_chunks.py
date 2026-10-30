#!/usr/bin/env python3
"""
Character-Level CTC Tokenization Pipeline for Audio Chunks.

Pipeline:
  1. Loads all 728 chunk transcripts from data/chunked_dataset/manifest_chunked.csv
  2. Builds complete multilingual CTC vocabulary for English, Gujarati, and Hindi.
  3. Saves vocabulary mapping to data/chunked_dataset/vocab.json
  4. Tokenizes all transcripts into character token IDs.
  5. Exports individual JSON token files to data/chunked_dataset/tokenized_transcripts/
  6. Exports updated manifest CSV to data/chunked_dataset/manifest_tokenized.csv
  7. Updates interactive Excel sheet at ~/Desktop/stt_chunked_summary.xlsx
"""

import csv
import json
import logging
import shutil
import sys
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from preprocessing.tokenizer import CTCTokenizer

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_BASE = PROJECT_ROOT / "data" / "chunked_dataset"
MANIFEST_CHUNKED = OUTPUT_BASE / "manifest_chunked.csv"
MANIFEST_TOKENIZED = OUTPUT_BASE / "manifest_tokenized.csv"
VOCAB_PATH = OUTPUT_BASE / "vocab.json"
TOKENIZED_DIR = OUTPUT_BASE / "tokenized_transcripts"
EXCEL_PATH = OUTPUT_BASE / "stt_chunked_summary.xlsx"
DESKTOP_EXCEL_PATH = Path.home() / "Desktop" / "stt_chunked_summary.xlsx"

LANG_NAME_MAP = {
    "en": "English 🇬🇧",
    "gu": "Gujarati 🇮🇳 (ગુજરાતી)",
    "hi": "Hindi 🇮🇳 (हिंदी)",
}

PORT = 8000


def main():
    logger.info("Starting Multilingual Character-Level CTC Tokenization Pipeline...")

    TOKENIZED_DIR.mkdir(parents=True, exist_ok=True)

    if not MANIFEST_CHUNKED.exists():
        logger.error(f"Manifest missing: {MANIFEST_CHUNKED}")
        return

    with open(MANIFEST_CHUNKED, "r", encoding="utf-8") as f:
        chunk_rows = list(csv.DictReader(f))

    logger.info(f"Loaded {len(chunk_rows)} chunk entries.")

    # Collect all transcript texts to build complete vocabulary
    all_texts = [r["transcript"] for r in chunk_rows if r.get("transcript")]

    # Build CTC Tokenizer
    tokenizer = CTCTokenizer()
    tokenizer.build_vocab(all_texts)
    tokenizer.save_vocab(VOCAB_PATH)

    logger.info(f"🔑 Vocabulary built successfully with {len(tokenizer)} tokens!")
    logger.info(f"📁 Vocab saved to: {VOCAB_PATH}")

    tokenized_entries = []

    for r in chunk_rows:
        audio_path = r["audio_path"]
        audio_name = Path(audio_path).name
        stem = Path(audio_path).stem
        transcript = r.get("transcript", "")

        # Encode text to CTC character token IDs
        token_ids = tokenizer.encode(transcript)
        token_chars = [tokenizer.id2char.get(tid, "") for tid in token_ids]

        # Save individual JSON token file
        json_path = TOKENIZED_DIR / f"{stem}.json"
        with open(json_path, "w", encoding="utf-8") as jf:
            json.dump({
                "audio_name": audio_name,
                "transcript": transcript,
                "token_ids": token_ids,
                "token_chars": token_chars,
                "num_tokens": len(token_ids),
                "language": r.get("language", "en"),
            }, jf, ensure_ascii=False, indent=2)

        entry = dict(r)
        entry["token_ids_str"] = str(token_ids)
        entry["token_ids"] = token_ids
        entry["num_tokens"] = len(token_ids)
        entry["json_path"] = str(json_path.resolve())
        tokenized_entries.append(entry)

    # Save manifest CSV with tokenization metadata
    fieldnames = ["audio_path", "original_source", "duration", "language", "transcript", "num_tokens", "token_ids_str"]
    with open(MANIFEST_TOKENIZED, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for e in tokenized_entries:
            writer.writerow({
                "audio_path": e["audio_path"],
                "original_source": e.get("original_source", ""),
                "duration": e["duration"],
                "language": e["language"],
                "transcript": e["transcript"],
                "num_tokens": e["num_tokens"],
                "token_ids_str": e["token_ids_str"],
            })

    logger.info(f"📄 Saved tokenized manifest CSV: {MANIFEST_TOKENIZED}")

    # Generate Updated Excel Sheet with Token Column
    tokenized_entries.sort(key=lambda x: x["audio_path"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2-5s Chunks & CTC Tokens"

    headers = [
        "S.No",
        "Play in Google Sheets",
        "Play Local File",
        "Chunk File Name",
        "Language",
        "Duration (sec)",
        "Accurate Native Transcript",
        "Token Count",
        "Character CTC Token IDs",
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    link_font_gs = Font(name="Calibri", size=11, color="008000", underline="single", bold=True)
    link_font_local = Font(name="Calibri", size=11, color="0000FF", underline="single")

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, e in enumerate(tokenized_entries, start=1):
        row_num = idx + 1
        audio_name = Path(e["audio_path"]).name
        lang = e["language"]
        lang_display = LANG_NAME_MAP.get(lang, lang)

        ws.cell(row=row_num, column=1, value=idx).alignment = Alignment(horizontal="center")

        gs_play_cell = ws.cell(row=row_num, column=2)
        gs_url = f"http://localhost:{PORT}/wav_chunks/{audio_name}"
        gs_play_cell.value = f'=HYPERLINK("{gs_url}", "▶️ Play in Browser / Sheets")'
        gs_play_cell.font = link_font_gs
        gs_play_cell.alignment = Alignment(horizontal="center")

        local_play_cell = ws.cell(row=row_num, column=3)
        local_url = f"file://{e['audio_path']}"
        local_play_cell.value = f'=HYPERLINK("{local_url}", "📁 Play Local File")'
        local_play_cell.font = link_font_local
        local_play_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=row_num, column=4, value=audio_name)
        ws.cell(row=row_num, column=5, value=lang_display)
        ws.cell(row=row_num, column=6, value=float(e["duration"])).alignment = Alignment(horizontal="right")
        ws.cell(row=row_num, column=7, value=e["transcript"])
        ws.cell(row=row_num, column=8, value=e["num_tokens"]).alignment = Alignment(horizontal="right")
        ws.cell(row=row_num, column=9, value=e["token_ids_str"])

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 70
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 65

    wb.save(EXCEL_PATH)
    shutil.copy(EXCEL_PATH, DESKTOP_EXCEL_PATH)

    logger.info(f"📊 Tokenized Excel summary saved to:")
    logger.info(f"  - {EXCEL_PATH}")
    logger.info(f"  - {DESKTOP_EXCEL_PATH}")
    logger.info("🎉 Character-Level CTC Tokenization completed successfully!")


if __name__ == "__main__":
    main()
